#!/usr/bin/env python3
"""Export every prompt and session id to docs/prompts-sessions.xlsx (needs: pip install openpyxl).

  scripts/export-sessions.py --since 2026-09-23 \
      --projects ~/.claude/projects/<repo-slug> ~/.claude/projects/-home-eslam-Downloads-TTFolder \
      [--include-sessions ids.txt] [--out docs/prompts-sessions.xlsx]

Sheets: Headless runs, Interactive sessions, Milestones, Summary, Human inputs.
Headless session ids (docs/sessions.csv) are left out of the interactive sheet.
It prints what it picked up, so you can check the sheet before committing it.
"""
import argparse
import csv
import json
import subprocess
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def read_csv(path, tag=None):
    if not path.exists():
        return []
    rows = list(csv.DictReader(path.open()))
    for r in rows:
        r["source"] = tag or "run"
    return rows


def num(v, cast=int):
    try:
        return cast(float(v))
    except (TypeError, ValueError):
        return 0


def commit_for(r):
    loop, feature, ms = r["loop"], r["feature"], r["milestone"]
    if loop == "testing" and r["mode"] == "phase":
        return ""
    if loop == "testing":
        grep = f"testing: {feature} suite {ms.split(':', 1)[-1]}"
    elif ms == "phase-00-plan":
        grep = f"{loop}: {feature} phase-00 plan"
    elif loop == "orchestrator":
        grep = f"orchestrator: {feature} plan"
    else:
        grep = f"{loop}: {feature} {ms} "
    out = subprocess.run(["git", "log", "--format=%h", "--fixed-strings", f"--grep={grep}"],
                         cwd=ROOT, capture_output=True, text=True).stdout.split()
    return out[0] if out else ""


def guess_step(text):
    t = text.lower()
    for step, words in (("constitution", ["constitution"]), ("setup", ["init", "install", "setup", "config"]),
                        ("debugging", ["error", "fail", "fix", "bug", "debug"]),
                        ("review", ["review", "approve"]), ("planning", ["plan", "prd", "spec"])):
        if any(w in t for w in words):
            return step
    return "building loops"


def prompt_text(msg):
    content = msg.get("content")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = [c.get("text", "") for c in content if isinstance(c, dict) and c.get("type") == "text"]
        return "\n".join(p for p in parts if p)
    return ""


def interactive(projects, since, include, headless_ids):
    out = []
    for proj in projects:
        for f in sorted(Path(proj).expanduser().glob("*.jsonl")):
            sid = f.stem
            if sid in headless_ids or (include and sid not in include):
                continue
            for line in f.open(errors="replace"):
                try:
                    e = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if e.get("type") != "user" or e.get("isMeta") or e.get("toolUseResult"):
                    continue
                text = prompt_text(e.get("message") or {})
                if not text.strip() or text.lstrip().startswith("<"):  # tool results / system notes
                    continue
                ts = e.get("timestamp", "")
                if since and ts and ts[:10] < since:
                    continue
                out.append({"session_id": e.get("sessionId", sid), "project": Path(proj).name,
                            "timestamp": ts, "step": guess_step(text), "prompt": text})
    return out


def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [min(len(str(r[i - 1])), 80) for r in rows if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 82)
    ws.freeze_panes = "A2"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--projects", nargs="*", default=[])
    ap.add_argument("--since", default="")
    ap.add_argument("--include-sessions")
    ap.add_argument("--out", default=str(DOCS / "prompts-sessions.xlsx"))
    a = ap.parse_args()

    runs = read_csv(DOCS / "sessions.csv") + read_csv(DOCS / "dry-run-sessions.csv", "dry-run")
    headless_ids = {r["session_id"] for r in runs if r["session_id"]}
    include = set(Path(a.include_sessions).read_text().split()) if a.include_sessions else set()
    inter = interactive(a.projects, a.since, include, headless_ids)
    humans = read_csv(DOCS / "human-inputs.csv") + read_csv(DOCS / "dry-run-human-inputs.csv", "dry-run")

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Headless runs",
          ["session_id", "source", "run_id", "engine_version", "started_at", "ended_at", "loop", "mode",
           "feature", "milestone", "kind", "attempt", "model", "phase_goal", "prompt", "context",
           "input tokens", "cache tokens", "output tokens", "cost USD", "duration ms", "result", "commit"],
          [[r["session_id"], r["source"], r["run_id"], r["engine_version"], r["started_at"], r["ended_at"],
            r["loop"], r["mode"], r["feature"], r["milestone"], r["kind"], num(r["attempt"]), r["model"],
            r["phase_goal"], r["prompt"], r["context"], num(r["input_tokens"]),
            num(r["cache_creation_input_tokens"]) + num(r["cache_read_input_tokens"]),
            num(r["output_tokens"]), num(r["total_cost_usd"], float), num(r["duration_ms"]), r["result"],
            commit_for(r) if r["source"] == "run" else ""] for r in runs])
    sheet(wb, "Interactive sessions", ["session_id", "project", "timestamp", "step", "prompt"],
          [[i["session_id"], i["project"], i["timestamp"], i["step"], i["prompt"]] for i in inter])

    ms = defaultdict(list)
    for r in runs:
        dev = r["milestone"].split(":", 1)
        key = (r["source"], r["feature"],
               f"{dev[0]}:{dev[1]}" if r["loop"] == "testing" and r["mode"] == "phase" else f"{r['loop']}:{r['milestone']}")
        ms[key].append(r)
    sheet(wb, "Milestones",
          ["source", "feature", "milestone", "sessions", "attempts", "input tokens", "cache tokens",
           "output tokens", "cost USD", "active ms", "last result"],
          [[k[0], k[1], k[2], len(v), max(num(r["attempt"]) for r in v), sum(num(r["input_tokens"]) for r in v),
            sum(num(r["cache_creation_input_tokens"]) + num(r["cache_read_input_tokens"]) for r in v),
            sum(num(r["output_tokens"]) for r in v), round(sum(num(r["total_cost_usd"], float) for r in v), 4),
            sum(num(r["duration_ms"]) for r in v), v[-1]["result"]] for k, v in ms.items()])

    once = {}
    for r in runs:
        once.setdefault(r["session_id"] or id(r), r)
    by = defaultdict(list)
    for r in once.values():
        by[(r["source"], r["loop"], r["mode"])].append(r)
    blocked = skipped = 0
    for f in ROOT.glob("loops/*/state/*/phases.json"):
        for p in json.loads(f.read_text()).get("phases", []):
            blocked += p["status"] == "blocked"
            skipped += p["status"] == "skipped"
    summ = [[k[0], k[1], k[2], len(v), sum(num(r["input_tokens"]) for r in v),
             sum(num(r["cache_creation_input_tokens"]) + num(r["cache_read_input_tokens"]) for r in v),
             sum(num(r["output_tokens"]) for r in v), round(sum(num(r["total_cost_usd"], float) for r in v), 4),
             sum(num(r["duration_ms"]) for r in v)] for k, v in sorted(by.items())]
    allr = list(once.values())
    summ.append(["TOTAL", "", "", len(allr), sum(num(r["input_tokens"]) for r in allr),
                 sum(num(r["cache_creation_input_tokens"]) + num(r["cache_read_input_tokens"]) for r in allr),
                 sum(num(r["output_tokens"]) for r in allr),
                 round(sum(num(r["total_cost_usd"], float) for r in allr), 4),
                 sum(num(r["duration_ms"]) for r in allr)])
    summ += [[], ["attempts (max per milestone, summed)", sum(max(num(r["attempt"]) for r in v) for v in ms.values())],
             ["blocked phases", blocked], ["skipped phases", skipped],
             ["interactive prompts", len(inter)], ["generated", datetime.now().isoformat(timespec="seconds")]]
    sheet(wb, "Summary", ["source", "loop", "mode", "sessions", "input tokens", "cache tokens",
                          "output tokens", "cost USD", "active ms"], summ)
    sheet(wb, "Human inputs", ["source", "at", "loop", "feature", "milestone", "gate", "decision", "notes",
                               "edit_diff", "read_by_session"],
          [[h["source"], h["at"], h["loop"], h["feature"], h["milestone"], h["gate"], h["decision"],
            h["notes"], h["edit_diff"], h["read_by_session"]] for h in humans])
    wb.save(a.out)
    print(f"headless sessions: {len(runs)} ({len(headless_ids)} ids)")
    print(f"interactive prompts: {len(inter)} from sessions:")
    for sid in sorted({i['session_id'] for i in inter}):
        print(f"  {sid}")
    print(f"human inputs: {len(humans)}")
    print(f"wrote {a.out}")


if __name__ == "__main__":
    main()
